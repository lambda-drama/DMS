# Copyright (c) 2026, Mania and contributors
"""Import Vehicle Models and Vehicle Service Items from FRT labour Excel workbooks."""

from __future__ import annotations

import os
import re

import frappe
from frappe import _
from frappe.utils import cint, flt, get_files_path

from dms.overrides.vehicle_service_item import (
	_erpnext_item_link_fieldname,
	ensure_labour_erpnext_item,
	sync_labour_erpnext_item_name,
)

SKIP_SHEETS = frozenset({"Remarks Sheet", "Sum of Models", "Sum"})
DEFAULT_BRAND = "JETOUR"
DEFAULT_FUEL_TYPE = "Petrol"
DEFAULT_TRANSMISSION = "Automatic (AT)"


class _SheetView:
	"""Minimal sheet API shared by xlrd and openpyxl loaders."""

	def __init__(self, rows: list[list]):
		self._rows = rows
		self.nrows = len(rows)
		self.ncols = max((len(r) for r in rows), default=0)

	def cell_value(self, row: int, col: int):
		if row < 0 or row >= self.nrows or col < 0 or col >= len(self._rows[row]):
			return ""
		value = self._rows[row][col]
		return "" if value is None else value


def _load_workbook_sheets(file_path: str) -> list[tuple[str, object]]:
	lower = (file_path or "").lower()
	if lower.endswith((".xlsx", ".xlsm")):
		try:
			from openpyxl import load_workbook
		except ImportError as exc:
			raise ImportError(_("Excel import requires openpyxl. Install it in the bench environment.")) from exc
		book = load_workbook(file_path, data_only=True, read_only=True)
		try:
			return [(name, _SheetView([list(row) for row in book[name].iter_rows(values_only=True)])) for name in book.sheetnames]
		finally:
			book.close()

	try:
		import xlrd
	except ImportError as exc:
		raise ImportError(_("Excel import requires xlrd. Install it in the bench environment.")) from exc
	book = xlrd.open_workbook(file_path)
	return [(name, book.sheet_by_name(name)) for name in book.sheet_names()]


def import_frt_workbook(file_path: str, brand: str = DEFAULT_BRAND) -> dict:
	"""Parse workbook and upsert Vehicle Models + Vehicle Service Items per sheet."""
	if not os.path.isfile(file_path):
		frappe.throw(_("File not found: {0}").format(file_path))

	ensure_brand(brand)
	sheets = _load_workbook_sheets(file_path)
	summary = {
		"models_created": 0,
		"models_updated": 0,
		"services_created": 0,
		"services_updated": 0,
		"services_skipped": 0,
		"sheets_processed": 0,
		"errors": [],
		"details": [],
	}

	for sheet_name, sheet in sheets:
		if sheet_name in SKIP_SHEETS:
			continue
		try:
			result = _import_model_sheet(sheet, sheet_name, brand)
			if not result:
				frappe.db.commit()
				continue
			summary["sheets_processed"] += 1
			summary["models_created"] += result["models_created"]
			summary["models_updated"] += result["models_updated"]
			summary["services_created"] += result["services_created"]
			summary["services_updated"] += result["services_updated"]
			summary["services_skipped"] += result["services_skipped"]
			summary["details"].append(result)
			frappe.db.commit()
		except Exception as exc:
			frappe.db.rollback()
			frappe.log_error(title=f"FRT import — {sheet_name}")
			summary["errors"].append({"sheet": sheet_name, "error": str(exc)})

	return summary


def import_frt_file_url(file_url: str, brand: str = DEFAULT_BRAND) -> dict:
	path = _resolve_file_path(file_url)
	return import_frt_workbook(path, brand=brand)


def _resolve_file_path(file_url: str) -> str:
	file_url = (file_url or "").strip()
	if not file_url:
		frappe.throw(_("File URL is required"))

	if file_url.startswith("/private/files/"):
		return get_files_path(*file_url.replace("/private/files/", "").split("/"), is_private=1)

	if file_url.startswith("/files/"):
		return get_files_path(*file_url.replace("/files/", "").split("/"), is_private=0)

	if os.path.isabs(file_url) and os.path.isfile(file_url):
		return file_url

	frappe.throw(_("Could not resolve uploaded file path"))


def ensure_brand(brand: str) -> str:
	brand = (brand or DEFAULT_BRAND).strip()
	if not brand:
		brand = DEFAULT_BRAND
	if not frappe.db.exists("Brand", brand):
		frappe.get_doc({"doctype": "Brand", "brand": brand}).insert(ignore_permissions=True)
	return brand


def ensure_vehicle_service_type(category: str) -> str:
	"""Ensure a Vehicle Service Type exists for the FRT category label."""
	category = (category or "").strip()
	if not category:
		return ""

	if frappe.db.exists("Vehicle Service Type", category):
		return category

	existing = frappe.db.get_value(
		"Vehicle Service Type", {"service_type_name": category}, "name"
	)
	if existing:
		return existing

	doc = frappe.get_doc(
		{
			"doctype": "Vehicle Service Type",
			"service_type_name": category,
			"is_active": 1,
		}
	)
	doc.insert(ignore_permissions=True)
	return doc.name


def _import_model_sheet(sheet, sheet_name: str, brand: str) -> dict | None:
	if sheet.nrows < 3:
		return None

	header_row_idx, colmap = _find_header_row(sheet)
	if not _is_frt_model_sheet(sheet_name, colmap):
		return None

	model_name, model_code, model_year = _sheet_model_meta(sheet, sheet_name, header_row_idx, colmap)
	if not model_code:
		return None

	vehicle_model = ensure_vehicle_model(
		model_code=model_code,
		model_name=model_name,
		brand=brand,
		model_year=model_year,
	)

	stats = {
		"sheet": sheet_name,
		"vehicle_model": vehicle_model,
		"model_code": model_code,
		"model_name": model_name,
		"models_created": 0,
		"models_updated": 0,
		"services_created": 0,
		"services_updated": 0,
		"services_skipped": 0,
	}

	for row_idx in range(header_row_idx + 1, sheet.nrows):
		row = [_cell(sheet, row_idx, col) for col in range(sheet.ncols)]
		description = _text(row, colmap.get("description"))
		if not description:
			continue

		category = _text(row, colmap.get("category"))
		hours = flt(row[colmap["hours"]]) if colmap.get("hours") is not None else 0
		row_model_code = _text(row, colmap.get("model_code")) or model_code
		cat_code = _text(row, colmap.get("cat_code"))
		sub_code_raw = row[colmap["sub_code"]] if colmap.get("sub_code") is not None else ""
		sub_code = _format_sub_code(sub_code_raw, cat_code)
		service_code = _normalize_service_code(_text(row, colmap.get("service_code")))

		if not service_code and row_model_code and cat_code and sub_code:
			service_code = _normalize_service_code(f"{row_model_code}{cat_code.upper()}{sub_code}")

		if not service_code:
			stats["services_skipped"] += 1
			continue

		created = upsert_vehicle_service_item(
			vehicle_model=vehicle_model,
			service_code=service_code,
			description=description,
			category=category,
			cat_code=cat_code,
			sub_code=sub_code,
			hours=hours,
			model_code=row_model_code,
		)
		if created:
			stats["services_created"] += 1
		else:
			stats["services_updated"] += 1

	return stats


def _parse_sheet_vehicle_model(sheet_name: str) -> tuple[str, str] | None:
	"""Parse tab titles like X50-JX50 into (model_name, model_code)."""
	sheet_name = (sheet_name or "").strip()
	if "-" not in sheet_name:
		return None

	left, right = sheet_name.split("-", 1)
	left = left.strip()
	right = right.strip()
	if not left or not _looks_like_model_code(right):
		return None

	return left, right.upper()


def _is_frt_model_sheet(sheet_name: str, colmap: dict) -> bool:
	"""Import any tab that looks like labour services for one vehicle model."""
	if colmap.get("description") is None:
		return False
	if colmap.get("service_code") is not None:
		return True
	return colmap.get("model_code") is not None and colmap.get("cat_code") is not None


def _first_row_model_code(sheet, header_row_idx: int, colmap: dict) -> str:
	idx = colmap.get("model_code")
	if idx is None:
		return ""
	for row_idx in range(header_row_idx + 1, sheet.nrows):
		code = _text([_cell(sheet, row_idx, col) for col in range(sheet.ncols)], idx)
		if code:
			return code.upper()
	return ""


def _sheet_model_meta(sheet, sheet_name: str, header_row_idx: int = 0, colmap: dict | None = None) -> tuple[str, str, int | None]:
	model_year = None
	parsed = _parse_sheet_vehicle_model(sheet_name)
	colmap = colmap or {}

	if parsed:
		model_name, model_code = parsed
	else:
		model_name = (sheet_name or "").strip()
		model_code = ""

	row_model_code = _first_row_model_code(sheet, header_row_idx, colmap)
	if row_model_code:
		model_code = row_model_code

	# Row 0 may carry the model year at the end of the sheet.
	if sheet.nrows:
		row0 = [_cell(sheet, 0, c) for c in range(sheet.ncols)]
		for val in reversed(row0):
			if _looks_numeric(val):
				year = cint(val)
				if 1980 <= year <= 2100:
					model_year = year
					break

	if not model_code:
		model_code = _slug_code(model_name)

	return model_name, model_code.upper(), model_year


def _find_header_row(sheet) -> tuple[int, dict]:
	for row_idx in range(min(5, sheet.nrows)):
		row = [_cell(sheet, row_idx, c) for c in range(sheet.ncols)]
		colmap = _header_map(row)
		if colmap.get("description") is not None and (
			colmap.get("hours") is not None or colmap.get("cat_code") is not None
		):
			return row_idx, colmap
	return 1, {}


def _header_map(row) -> dict:
	colmap: dict[str, int] = {}
	# Workbooks duplicate headers on the right — only read the left block.
	limit = min(len(row), 8)
	for idx in range(limit):
		raw = row[idx]
		key = _text_value(raw).lower()
		if not key:
			continue
		if "job description" in key:
			colmap.setdefault("description", idx)
		elif key == "category":
			colmap.setdefault("category", idx)
		elif key in ("hours", "frt", "front"):
			colmap.setdefault("hours", idx)
		elif key == "model" or "model code" in key:
			colmap.setdefault("model_code", idx)
		elif "cat code" in key:
			colmap.setdefault("cat_code", idx)
		elif "sub code" in key:
			colmap.setdefault("sub_code", idx)
		elif "service" in key and "code" in key:
			colmap.setdefault("service_code", idx)
	return colmap


def _existing_item_name(item_code: str) -> str | None:
	item_code = (item_code or "").strip()
	if not item_code:
		return None
	return frappe.db.exists("Item", item_code) or frappe.db.get_value(
		"Item", {"item_code": item_code}, "name"
	)


def ensure_vehicle_model(
	model_code: str,
	model_name: str,
	brand: str,
	model_year: int | None = None,
) -> str:
	model_code = (model_code or "").strip().upper()
	model_name = (model_name or model_code).strip()
	if not model_code:
		frappe.throw(_("Model code is required"))

	existing = (
		frappe.db.get_value("Vehicle Model", {"model_code": model_code}, "name")
		or frappe.db.exists("Vehicle Model", model_code)
	)
	if existing:
		_update_vehicle_model_fields(existing, model_name, brand, model_year)
		return existing

	item_code = ensure_vehicle_item(model_code, model_name, brand)
	existing = frappe.db.exists("Vehicle Model", item_code) or frappe.db.get_value(
		"Vehicle Model", {"model": item_code}, "name"
	)
	if existing:
		_update_vehicle_model_fields(existing, model_name, brand, model_year)
		return existing

	doc = frappe.get_doc(
		{
			"doctype": "Vehicle Model",
			"model": item_code,
			"model_code": model_code,
			"model_name": model_name,
			"brand": brand,
			"model_year": model_year,
			"fuel_type": DEFAULT_FUEL_TYPE,
			"transmission": DEFAULT_TRANSMISSION,
			"is_active": 1,
		}
	)
	try:
		doc.insert(ignore_permissions=True, ignore_if_duplicate=True)
	except frappe.DuplicateEntryError:
		existing = (
			frappe.db.exists("Vehicle Model", item_code)
			or frappe.db.get_value("Vehicle Model", {"model_code": model_code}, "name")
			or item_code
		)
		_update_vehicle_model_fields(existing, model_name, brand, model_year)
		return existing
	return doc.name


def _update_vehicle_model_fields(
	name: str, model_name: str, brand: str, model_year: int | None = None
) -> None:
	"""Write model fields without get_doc/save (avoids TimestampMismatchError)."""
	current = frappe.db.get_value(
		"Vehicle Model",
		name,
		["model_name", "brand", "is_active", "model_year", "model_code"],
		as_dict=True,
	)
	if not current:
		return

	updates = {}
	if (current.get("model_name") or "") != model_name:
		updates["model_name"] = model_name
	if (current.get("brand") or "") != brand:
		updates["brand"] = brand
	if not cint(current.get("is_active")):
		updates["is_active"] = 1
	if model_year and cint(current.get("model_year")) != cint(model_year):
		updates["model_year"] = model_year
	if not updates:
		return

	frappe.db.set_value("Vehicle Model", name, updates, update_modified=False)


def ensure_vehicle_item(item_code: str, item_name: str, brand: str) -> str:
	item_code = (item_code or "").strip().upper()
	existing = _existing_item_name(item_code)
	if existing:
		return existing

	item_group = _vehicle_item_group()
	try:
		frappe.get_doc(
			{
				"doctype": "Item",
				"item_code": item_code,
				"item_name": item_name[:140] or item_code,
				"item_group": item_group,
				"brand": brand,
				"stock_uom": "Nos",
				"is_stock_item": 1,
				"has_serial_no": 1,
				"is_sales_item": 1,
				"is_purchase_item": 1,
			}
		).insert(ignore_permissions=True, ignore_if_duplicate=True)
	except frappe.DuplicateEntryError:
		return _existing_item_name(item_code) or item_code

	return _existing_item_name(item_code) or item_code


def _vehicle_item_group() -> str:
	groups = frappe.get_all("Item Group", filters={"custom_is_vehicle": 1}, pluck="name", limit=1)
	if groups:
		return groups[0]
	if frappe.db.exists("Item Group", "Vehicles"):
		return "Vehicles"
	return frappe.db.get_single_value("DMS Settings", "default_item_group") or "All Item Groups"


def upsert_vehicle_service_item(
	vehicle_model: str,
	service_code: str,
	description: str,
	category: str,
	cat_code: str,
	sub_code: str,
	hours: float,
	model_code: str,
) -> bool:
	service_code = _normalize_service_code(service_code)
	description = (description or service_code).strip()
	if not service_code:
		return False

	meta = frappe.get_meta("Vehicle Service Item")
	link_field = _erpnext_item_link_fieldname()

	existing_name = frappe.db.get_value(
		"Vehicle Service Item",
		{"custom_service_code": service_code, "custom_vehicle_model": vehicle_model},
		"name",
	)
	if not existing_name and frappe.db.exists("Vehicle Service Item", service_code):
		candidate = frappe.db.get_value("Vehicle Service Item", service_code, "custom_vehicle_model")
		if candidate == vehicle_model:
			existing_name = service_code

	category_link = ensure_vehicle_service_type(category)
	service_item_name = _service_item_display_name(description, service_code, existing_name)

	values = {
		"service_item": service_item_name,
		"custom_item_name": description,
		"custom_service_code": service_code,
		"custom_vehicle_model": vehicle_model,
		"custom_category": category_link,
		"custom_cat_code": cat_code,
		"custom_sub_code": sub_code,
		"custom_frt": _hours_text(hours),
	}

	if meta.has_field("custom_estimated_timehours"):
		values["custom_estimated_timehours"] = _hours_text(hours)
	if meta.has_field("custom_active"):
		values["custom_active"] = 1

	if existing_name:
		doc = frappe.get_doc("Vehicle Service Item", existing_name)
		doc.flags.ignore_if_modified = True
		doc.update(values)
		if not (doc.get(link_field) or "").strip():
			ensure_labour_erpnext_item(doc, link_field)
		else:
			sync_labour_erpnext_item_name(doc, link_field)
		doc.save(ignore_permissions=True)
		return False

	doc = frappe.get_doc({"doctype": "Vehicle Service Item", **values})
	ensure_labour_erpnext_item(doc, link_field)
	doc.insert(ignore_permissions=True)
	return True


def _service_item_display_name(
	description: str, service_code: str, existing_name: str | None = None
) -> str:
	"""Use job description as Service Item name; suffix code if needed for uniqueness."""
	description = (description or service_code).strip()
	if not description:
		return service_code

	if not frappe.db.exists("Vehicle Service Item", description):
		return description

	owner = frappe.db.get_value("Vehicle Service Item", description, "name")
	if existing_name and owner == existing_name:
		return description

	fallback = f"{description} ({service_code})"
	if not frappe.db.exists("Vehicle Service Item", fallback):
		return fallback

	return service_code


def _normalize_service_code(value: str) -> str:
	return re.sub(r"\s+", "", (value or "")).upper()


def _hours_text(hours: float) -> str:
	if not hours:
		return "0"
	if hours == int(hours):
		return str(int(hours))
	return str(round(hours, 2))


def _format_sub_code(value, cat_code: str) -> str:
	if value in (None, ""):
		return ""
	cat = (cat_code or "").strip().upper()
	if isinstance(value, float):
		if value == int(value):
			n = int(value)
			return f"{n:02d}" if cat == "EN" else str(n)
		return str(value).strip()
	text = str(value).strip()
	if cat == "EN" and text.isdigit():
		return text.zfill(2)
	return text


def _cell(sheet, row: int, col: int):
	try:
		return sheet.cell_value(row, col)
	except IndexError:
		return ""


def _text(row, idx) -> str:
	if idx is None:
		return ""
	return _text_value(row[idx])


def _text_value(value) -> str:
	if value is None:
		return ""
	if isinstance(value, float):
		if value == int(value):
			return str(int(value))
		return str(value).strip()
	return str(value).strip()


def _looks_numeric(value) -> bool:
	try:
		float(value)
		return True
	except (TypeError, ValueError):
		return False


def _looks_like_model_code(code: str) -> bool:
	code = (code or "").strip()
	if len(code) < 3:
		return False
	return any(c.isalpha() for c in code)


def _slug_code(text: str) -> str:
	base = re.sub(r"[^\w\-]+", "-", (text or "").upper()).strip("-")
	return base or "MODEL"
